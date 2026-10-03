from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[3]
OUT = ROOT / "output" / "documents"
ASSET = ROOT / "assets" / "branding"
WORK = Path(__file__).resolve().parent
OUT.mkdir(parents=True, exist_ok=True)
WORK.mkdir(parents=True, exist_ok=True)
DOCX = OUT / "MFMA_Official_Sporting_Code_2026_Final_Draft.docx"

RED = "E31624"
BLACK = "111111"
DARK = "191D24"
MID = "5B6470"
LIGHT = "F2F4F7"
BORDER = "D9D9D9"


def font(size=20, bold=False):
    try:
        return ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial.ttf", size=size)
    except Exception:
        return ImageFont.load_default()


def brand_banner(source, target):
    logo = Image.open(source).convert("RGBA")
    logo.thumbnail((1500, 500), Image.Resampling.LANCZOS)
    canvas = Image.new("RGBA", (1600, 560), (10, 12, 16, 255))
    x = (canvas.width - logo.width) // 2
    y = (canvas.height - logo.height) // 2
    canvas.alpha_composite(logo, (x, y))
    canvas.convert("RGB").save(target, quality=95)


def arrow(draw, start, end, color=(30, 30, 30), width=7):
    draw.line([start, end], fill=color, width=width)
    x1, y1 = start; x2, y2 = end
    import math
    ang = math.atan2(y2-y1, x2-x1)
    for delta in (2.55, -2.55):
        draw.line([end, (x2+18*math.cos(ang+delta), y2+18*math.sin(ang+delta))], fill=color, width=width)


def flag_polygon(draw, pole_x, pole_y, color, direction="right", folded=False, checked=False):
    draw.line([(pole_x, pole_y), (pole_x, pole_y+150)], fill=(35,35,35), width=8)
    if direction == "right":
        pts = [(pole_x+4,pole_y+8),(pole_x+132,pole_y+18),(pole_x+112,pole_y+91),(pole_x+4,pole_y+80)]
    elif direction == "left":
        pts = [(pole_x-4,pole_y+8),(pole_x-132,pole_y+18),(pole_x-112,pole_y+91),(pole_x-4,pole_y+80)]
    elif direction == "up-right":
        pts = [(pole_x+2,pole_y+5),(pole_x+100,pole_y-75),(pole_x+130,pole_y-5),(pole_x+20,pole_y+55)]
    else:
        pts = [(pole_x-2,pole_y+5),(pole_x-100,pole_y-75),(pole_x-130,pole_y-5),(pole_x-20,pole_y+55)]
    if folded:
        if direction in ("right","left"):
            sign = 1 if direction == "right" else -1
            pts = [(pole_x+sign*3,pole_y+10),(pole_x+sign*65,pole_y+22),(pole_x+sign*55,pole_y+42),(pole_x+sign*3,pole_y+35)]
        else:
            sign = 1 if direction == "up-right" else -1
            pts = [(pole_x,pole_y+5),(pole_x+sign*48,pole_y-38),(pole_x+sign*62,pole_y-24),(pole_x+sign*12,pole_y+20)]
    if checked:
        draw.polygon(pts, fill=(245,245,245), outline=(20,20,20))
        minx,miny,maxx,maxy = min(p[0] for p in pts),min(p[1] for p in pts),max(p[0] for p in pts),max(p[1] for p in pts)
        step=22
        for yy in range(miny,maxy,step):
            for xx in range(minx,maxx,step):
                if ((xx-minx)//step+(yy-miny)//step)%2==0:
                    draw.rectangle([xx,yy,min(xx+step,maxx),min(yy+step,maxy)], fill=(15,15,15))
        draw.line(pts+[pts[0]], fill=(20,20,20), width=2)
    else:
        draw.polygon(pts, fill=color, outline=(20,20,20))


def make_signal_diagram(name, title):
    img = Image.new("RGB", (700, 300), "white")
    d = ImageDraw.Draw(img)
    d.ellipse([92,50,135,93], fill=(35,35,35))
    d.line([(114,93),(114,218)], fill=(35,35,35), width=10)
    d.line([(114,130),(65,187)], fill=(35,35,35), width=9)
    d.line([(114,130),(180,150)], fill=(35,35,35), width=9)
    d.line([(114,218),(75,276)], fill=(35,35,35), width=9)
    d.line([(114,218),(152,276)], fill=(35,35,35), width=9)
    if name == "green":
        flag_polygon(d,180,68,(15,158,63))
    elif name == "safety-car":
        flag_polygon(d,180,68,(255,226,0))
        d.line([(180,145),(355,145)], fill=(35,35,35), width=8)
        d.polygon([(225,145),(245,48),(322,60),(302,145)], fill=(208,0,0), outline=(20,20,20))
        d.text((382,95),"YELLOW OPEN\nRED PERPENDICULAR",font=font(22,True),fill=(20,20,20),spacing=7)
    elif name == "investigation":
        flag_polygon(d,180,68,(248,248,248))
    elif name == "warning":
        flag_polygon(d,180,68,(248,248,248))
        flag_polygon(d,180,175,(255,226,0),direction="up-left",folded=True)
        d.text((382,95),"WHITE OPEN\nYELLOW FOLDED",font=font(22,True),fill=(20,20,20),spacing=7)
    elif name == "return":
        flag_polygon(d,180,105,(255,226,0),folded=True)
        d.arc([185,38,330,195], start=205, end=520, fill=(45,45,45), width=7)
        arrow(d,(298,166),(320,139))
        d.text((382,95),"TWIRL BRIEFLY\nSTOP AND REPEAT",font=font(22,True),fill=(20,20,20),spacing=7)
    elif name == "proceed":
        flag_polygon(d,180,105,(15,158,63),folded=True)
        arrow(d,(210,132),(355,132))
        d.text((382,95),"TWIRL THEN POINT\nTO STARTING LINE",font=font(22,True),fill=(20,20,20),spacing=7)
    elif name == "halfway":
        flag_polygon(d,175,135,(15,158,63),direction="up-right")
        flag_polygon(d,175,135,(255,255,255),direction="up-left",checked=True)
        d.text((382,95),"GREEN CROSSED\nWITH CHECKERED",font=font(22,True),fill=(20,20,20),spacing=7)
    d.text((22,12),title,font=font(24,True),fill=(10,10,10))
    path = WORK / f"signal_{name}.png"
    img.save(path)
    return path


banner = WORK / "mra_banner.jpg"
brand_banner(ASSET / "mra-logo.png", banner)
diagrams = {k:make_signal_diagram(k,v) for k,v in {
    "green":"START OR RESUME","safety-car":"SAFETY CAR","investigation":"INVESTIGATION",
    "warning":"INFRACTION WARNING","return":"RETURN TO START","proceed":"PROCEED TO LINE",
    "halfway":"QUALIFYING HALFWAY"
}.items()}

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5); sec.page_height = Inches(11)
sec.top_margin = Inches(.68); sec.bottom_margin = Inches(.65)
sec.left_margin = Inches(.75); sec.right_margin = Inches(.75)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"; normal._element.rPr.rFonts.set(qn("w:ascii"),"Arial"); normal._element.rPr.rFonts.set(qn("w:hAnsi"),"Arial"); normal.font.size = Pt(10.5); normal.font.color.rgb = RGBColor.from_string(BLACK)
normal.paragraph_format.space_after = Pt(5); normal.paragraph_format.line_spacing = 1.08
for name,size,bold,space in [("Title",30,True,14),("Subtitle",13,False,10),("Heading 1",19,True,12),("Heading 2",13,True,7),("Heading 3",11,True,5)]:
    st=styles[name]; st.font.name="Arial"; st._element.rPr.rFonts.set(qn("w:ascii"),"Arial"); st._element.rPr.rFonts.set(qn("w:hAnsi"),"Arial"); st.font.size=Pt(size); st.font.bold=bold; st.font.color.rgb=RGBColor(0,0,0); st.paragraph_format.space_before=Pt(space); st.paragraph_format.space_after=Pt(5); st.paragraph_format.keep_with_next=True


def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn("w:shd"))
    if shd is None: shd=OxmlElement("w:shd"); tcPr.append(shd)
    shd.set(qn("w:fill"),fill)


def borders(table):
    tblPr=table._tbl.tblPr; tb=tblPr.find(qn("w:tblBorders"))
    if tb is None: tb=OxmlElement("w:tblBorders"); tblPr.append(tb)
    for edge in ("top","left","bottom","right","insideH","insideV"):
        e=OxmlElement(f"w:{edge}"); e.set(qn("w:val"),"single"); e.set(qn("w:sz"),"5"); e.set(qn("w:color"),BORDER); tb.append(e)


def cell_margin(cell, top=85, start=105, bottom=85, end=105):
    tc=cell._tc; tcPr=tc.get_or_add_tcPr(); mar=tcPr.first_child_found_in("w:tcMar")
    if mar is None: mar=OxmlElement("w:tcMar"); tcPr.append(mar)
    for m,v in (("top",top),("start",start),("bottom",bottom),("end",end)):
        node=mar.find(qn(f"w:{m}"))
        if node is None: node=OxmlElement(f"w:{m}"); mar.append(node)
        node.set(qn("w:w"),str(v)); node.set(qn("w:type"),"dxa")


def add_table(headers, rows, widths=None, font_size=9.2):
    table=doc.add_table(rows=1,cols=len(headers)); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.autofit=False
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for i,h in enumerate(headers):
        c=table.rows[0].cells[i]; shade(c,DARK); c.vertical_alignment=WD_ALIGN_VERTICAL.CENTER; cell_margin(c)
        p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(h); r.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(9)
        if widths: c.width=Inches(widths[i])
    for ridx,row in enumerate(rows):
        cells=table.add_row().cells
        for i,val in enumerate(row):
            c=cells[i]; c.vertical_alignment=WD_ALIGN_VERTICAL.CENTER; cell_margin(c)
            if ridx%2: shade(c,"F7F8FA")
            p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.alignment=WD_ALIGN_PARAGRAPH.CENTER if i==0 else WD_ALIGN_PARAGRAPH.LEFT
            r=p.add_run(str(val)); r.font.size=Pt(font_size)
            if widths: c.width=Inches(widths[i])
    borders(table)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return table


def add_page_number(paragraph):
    paragraph.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    run=paragraph.add_run("MFMA Official Sporting Code   |   ")
    run.font.size=Pt(8); run.font.color.rgb=RGBColor.from_string(MID)
    fldChar1=OxmlElement("w:fldChar"); fldChar1.set(qn("w:fldCharType"),"begin")
    instr=OxmlElement("w:instrText"); instr.set(qn("xml:space"),"preserve"); instr.text=" PAGE "
    fldChar2=OxmlElement("w:fldChar"); fldChar2.set(qn("w:fldCharType"),"end")
    run._r.append(fldChar1); run._r.append(instr); run._r.append(fldChar2)


add_page_number(sec.footer.paragraphs[0])


def title(text, level=1):
    return doc.add_heading(text, level=level)


def clause(number, text):
    p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.18); p.paragraph_format.first_line_indent=Inches(-.18); p.paragraph_format.keep_together=True
    r=p.add_run(number+"  "); r.bold=True
    p.add_run(text)
    return p


def article(number, name, sections):
    doc.add_page_break()
    title(f"Article {number} {name}",1)
    for secnum, secname, clauses in sections:
        title(f"{number}.{secnum} {secname}",2)
        for idx,text in enumerate(clauses,1): clause(f"{number}.{secnum}.{idx}",text)


def add_accessible_picture(paragraph, path, width, description):
    shape = paragraph.add_run().add_picture(str(path), width=width)
    shape._inline.docPr.set("descr", description)
    shape._inline.docPr.set("title", description)
    return shape


# Cover
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; add_accessible_picture(p, banner, Inches(6.75), "MFMA Race Authority logo in red, white, and black")
p=doc.add_paragraph(style="Title"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run("MFMA Official Sporting Code")
p=doc.add_paragraph(style="Subtitle"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run("2026 Final Draft")
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(16)
r=p.add_run("Maeschen Farm Motorsports Association"); r.bold=True; r.font.size=Pt(14)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run("Revised October 3 2026\nEffective upon adoption by the MFMA")
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(36)
r=p.add_run("Safety takes priority over competition at all times."); r.bold=True; r.font.size=Pt(13); r.font.color.rgb=RGBColor.from_string(RED)

doc.add_page_break()
title("Competition Control Quick Reference",1)
doc.add_paragraph("This sheet summarizes the most frequently used controls. The numbered provisions of this Sporting Code govern whenever a summary and a rule appear to conflict.")
add_table(["Signal","Meaning","Required action"],[
    ("Green","Start or resume","Normal competition begins or resumes. Maximum speed is 15 mph unless reduced by the MRA Steward."),
    ("Yellow","Global caution","Competition continues at no more than 7 mph. The official clock advances at one-half rate."),
    ("Yellow plus Red","Safety Car","Follow the Official Vehicle. No overtaking or competitive activity."),
    ("Red","Stop","Stop safely as soon as practical and await instructions."),
    ("White","Investigation","Return to the Starting Zone. No penalty has yet been decided."),
    ("White plus folded Yellow","Infraction warning","Continue with caution. The warning lasts 10 seconds."),
    ("Checkered","Session complete","Cease competition and follow the post-session movement order."),
], [1.25,1.7,3.85])
doc.add_paragraph("Official communications. Flags, MRA digital displays, and Official Radio Channel announcements are parts of one race-control system. A Red instruction always controls. Drivers must obey the safest applicable instruction.")
doc.add_paragraph("Standard regular session. One-minute Hiding Period followed by a two-minute Finding Period unless the MRA Steward announces another duration before the start.")
doc.add_paragraph("Spot. A spot is valid when visual identification occurs and the confirming radio transmission begins before official time expires.")

doc.add_page_break()
title("Contents",1)
for item in [
    "Article 1 General Provisions","Article 2 Definitions","Article 3 Governance and MRA Steward Authority","Article 4 Registration Teams and Vehicles",
    "Article 5 Event Administration","Article 6 Regular Event Competition","Article 7 Qualifying","Article 8 Sprint Operations",
    "Article 9 Championship and Scoring","Article 10 Communications and Digital Race Control","Article 11 Conduct Investigations and Penalties",
    "Article 12 Safety Course and Hazard Management","Article 13 Official Signals","Appendix A Course and Sector Administration",
    "Appendix B Physical Flag Reference","Appendix C Championship Tables","Appendix D Revision Record"
]:
    p=doc.add_paragraph(item); p.style=styles["Normal"]; p.paragraph_format.space_after=Pt(7)

article(1,"General Provisions",[
 (1,"Name and Authority",[
  "The Maeschen Farm Motorsports Association, abbreviated MFMA, is the governing authority for competitions conducted under this Sporting Code.",
  "This Sporting Code governs every MFMA-sanctioned event unless a provision expressly states otherwise.",
  "Participation constitutes acceptance of this Sporting Code, official signals, event briefings, and lawful temporary instructions issued by the MRA Steward."
 ]),
 (2,"Safety and Fairness",[
  "Safety takes priority over competition, scoring, procedure, schedule, and competitive advantage at all times.",
  "Nothing in this Sporting Code authorizes unsafe, reckless, unlawful, or negligent conduct.",
  "Competitors shall not exploit an omission, technicality, or system limitation to obtain an advantage inconsistent with fair competition."
 ]),
 (3,"Interpretation",[
  "The MRA Steward interprets and applies this Sporting Code during an event.",
  "When a rule is silent or incomplete, the MRA Steward may issue a temporary ruling consistent with safety, fairness, proportional competition, and event integrity.",
  "A temporary ruling applies to that event unless later incorporated into an adopted edition of this Sporting Code."
 ])
])

article(2,"Definitions",[
 (1,"Competition Terms",[
  "Event means the complete MFMA-sanctioned activity conducted under one event record, including any sessions, qualifying, Sprint Operations, or non-points activity.",
  "Regular Event means an event composed of scored sessions in which teams alternate Pursuit and Evading roles.",
  "Session means one individual Pursuit-versus-Evading competition within a Regular Event.",
  "Circuit means the minimum balanced group of sessions in which every participating team receives proportionate Pursuit and Evading opportunity.",
  "Qualifying means a Shelly time-trial competition on a Short, Medium, or Long approved course.",
  "Sprint Operations means competition or organized driving in which the MRA signal system is used without a formal Regular Event session order.",
  "Points Event means an event designated before its start as eligible to award championship points."
 ]),
 (2,"People and Organizations",[
  "MRA Steward means the official appointed as Race Director for the event. Race Director and MRA Steward refer to the same event authority unless additional stewards are expressly appointed.",
  "Competitor means a registered driver, passenger, or participating team member.",
  "Driver means the registered competitor operating a vehicle at that time.",
  "Passenger means a registered competitor riding in a competition vehicle without operating it.",
  "Team Representative means the one person designated by a team to submit rule questions, formal clarifications, or team-level representations to the MRA Steward."
 ]),
 (3,"Signals and Systems",[
  "Official Signal means a physical flag, MRA-authorized digital display, official sound or voice alert, or Official Radio Channel instruction.",
  "Primary Safety State means Green, Yellow, Safety Car, or Red.",
  "Supplemental Signal means a local Yellow, Grip warning, Move Over instruction, enforcement indication, or movement order that does not weaken the Primary Safety State.",
  "Official Vehicle means the vehicle assigned to the MRA Steward for race control, safety operations, and official flag display."
 ])
])

article(3,"Governance and MRA Steward Authority",[
 (1,"Appointment and Responsibility",[
  "Every sanctioned event shall operate under one MRA Steward.",
  "The MRA Steward administers the Sporting Code, establishes boundaries and formats, maintains official time, controls official signals, determines results, resolves safety reports, and records penalties.",
  "The MRA Steward may suspend, neutralize, restart, terminate, invalidate, or modify competition when necessary for safety, fairness, property protection, or event integrity."
 ]),
 (2,"Dual Participation",[
  "The MRA Steward may also compete when necessary, but official duties always take precedence.",
  "A dual-role Steward shall act impartially and shall preserve an auditable event record for decisions that affect the Steward's own team.",
  "Other teams may use their designated Team Representative for questions and clarification, but no Team Representative exercises race-control authority."
 ]),
 (3,"Official Decisions",[
  "Active-competition decisions are final for conducting the event.",
  "The MRA Steward should explain significant rulings when practical and shall record investigations, penalties, result adjustments, and event invalidations in the MRA system.",
  "A completed result may be corrected for a scoring, timing, registration, or administrative error, or through an authorized audited points adjustment."
 ]),
 (4,"Rule Amendments and Temporary Suspensions",[
  "A mid-season amendment requires a written motion identifying the affected rule, exact proposed change, reason, and effective time. The MRA Steward shall conduct and record a roll-call vote of one designated Team Representative for every recognized team. Adoption requires a Yes vote from every recognized team and MRA Steward authorization; a No vote, abstention, absence, or missing response is not consent.",
  "A temporary mid-session suspension may affect only a non-safety procedural or competition rule. Competition shall first be brought to a safe, paused state. The motion shall identify the exact rule, scope, duration, and reason. Adoption requires a Yes vote from every active recognized team and MRA Steward authorization. It applies only to the current session and expires automatically when that session ends unless ended earlier by the Steward.",
  "No vote may suspend or weaken a safety rule, official signal, stop or speed requirement, hazard-reporting duty, property or legal obligation, or MRA emergency authority. The Steward may act immediately for safety without a vote but shall record the action when practical.",
  "The official roll-call record shall state each team, representative name, Yes, No, or Abstain vote, Steward identity, timestamp, motion, result, and any temporary scope. No amendment may retroactively alter a completed result, earned points, or penalty."
 ])
])

article(4,"Registration Teams and Vehicles",[
 (1,"Driver Registration",[
  "A driver shall be registered under the driver's name, MRA number, status, and team affiliation before earning points.",
  "A team device may sign in to the Driver Portal and switch among approved team drivers without creating separate device accounts.",
  "An honorary participant may appear in event records but is not a scoring driver unless the MRA changes that participant's competition role before the relevant event."
 ]),
 (2,"Teams and Representatives",[
  "The championship currently recognizes the approved Monarch MFMA Team and Parakeet MFMA Team. The MFMA may approve additional teams by updating the official registry.",
  "Each team may designate one Team Representative for rules communication. Driver hazard reports remain available independently for immediate safety needs.",
  "Team championship points follow the driver's registered affiliation. A passenger's affiliation does not transfer the driver's points."
 ]),
 (3,"Vehicles",[
  "Ranger and Shelly are standard approved vehicles. Gator is an emergency-only vehicle and requires a recorded reason whenever selected.",
  "A newly submitted vehicle may participate only after MRA approval and addition to the official vehicle registry.",
  "An event vehicle replacement requires MRA approval and a recorded reason before a Points Event may be archived.",
  "Every vehicle shall be maintained in safe working condition and may be excluded until a reported defect is resolved."
 ])
])

article(5,"Event Administration",[
 (1,"Event Opening",[
  "Before opening an event, the MRA Steward shall select Regular Event, Sprint, or Qualifying and shall designate the event as points or non-points.",
  "The competitor briefing shall identify boundaries, open and closed sectors, vehicle assignments, format, timing, signals, Official Radio Channel, special conditions, and whether qualifying serves as a tiebreaker.",
  "Test Mode is solely for display, sound, and system verification. It creates no competition result, point award, hazard record, or official session."
 ]),
 (2,"Starting and Post-Session Procedure",[
  "All formal sessions begin at the Starting Zone unless the MRA Steward directs otherwise.",
  "The preparatory start board illuminates five paired groups of Red lights. The team shall be stopped behind the line before the sequence begins.",
  "When the lights extinguish, the Green signal is displayed. Official timing begins with Green, not with the preparatory lights.",
  "After Checkered, the Return to Start order is given by the folded Yellow movement signal and digital instruction.",
  "After competitors return and post-session matters are complete, the Proceed to Line order is given by the folded Green movement signal and digital instruction.",
  "An investigation may be opened only before Proceed to Line is issued."
 ]),
 (3,"Event Closure",[
  "The MRA Steward may end and archive a legal event, or end and invalidate an event at any time.",
  "A Points Event shall not be archived while roles are unbalanced, no official session exists, a safety report or investigation remains unresolved, an emergency replacement lacks approval, or no single final event winner has been recorded.",
  "An invalidated or void event remains an administrative record but awards no championship points."
 ])
])

article(6,"Regular Event Competition",[
 (1,"Roles and Timing",[
  "A Regular Event alternates Pursuit and Evading roles so each team receives proportionate opportunity.",
  "The standard Hiding Period is one minute and the standard Finding Period is two minutes. The MRA Steward may announce different times before the session begins.",
  "The official session clock advances at normal rate under Green, one-half rate under Yellow, and pauses under Red unless the MRA Steward announces a correction."
 ]),
 (2,"Spot and Result",[
  "A spot occurs when the Pursuit Team visually identifies the Evading Team and begins the confirming radio transmission before official time expires.",
  "The Pursuit Team wins by completing the required spot during the Finding Period. The Evading Team wins when required time expires without a valid spot.",
  "When multiple Pursuit Vehicles are formally required, the briefing shall state whether each must complete a spot.",
  "Only the MRA Steward may finalize a session result."
 ]),
 (3,"Restart and Neutralization",[
  "After a stopped or compromised session, the MRA Steward may order all competitors back to the line and initiate the full Red-light start sequence.",
  "A Safety Car neutralizes competitive activity. The MRA Steward may resume with Green, order a restart, or terminate the session after evaluating the circumstances.",
  "A Red condition pauses competition and requires a safe stop. A subsequent instruction shall state whether the session resumes, restarts, or terminates."
 ])
])

article(7,"Qualifying",[
 (1,"Format",[
  "Qualifying is conducted with Shelly on an approved Short, Medium, or Long course.",
  "Each driver receives a warm-up lap followed by no fewer than two and no more than ten timed laps, as established before qualifying begins.",
  "The Green-and-Checkered crossed signal marks the halfway point of the allotted timed laps.",
  "The fastest valid timed lap determines the qualifying winner. Invalidated laps remain in the audit record but are excluded from ranking."
 ]),
 (2,"Scoring and Tiebreaking",[
  "Standalone Points Qualifying awards one participation point to each registered driver who records a valid timed lap and one additional point to the fastest driver.",
  "Qualifying embedded in a Regular Event as a tiebreaker awards no separate qualifying points.",
  "When a Regular Event finishes tied and an official qualifying winner exists, the qualifying result determines the event winner unless the briefing established another approved procedure."
 ])
])

article(8,"Sprint Operations",[
 (1,"Purpose and Control",[
  "Sprint Operations permit organized competition and signal use without a formal Regular Event session order or finishing classification.",
  "The MRA Steward retains full safety and signal authority during Sprint Operations.",
  "A Sprint participant must be placed on the pre-start roster and complete at least one verified minute to qualify for a participation award."
 ]),
 (2,"Scoring",[
  "A Points Sprint awards one driver participation point per event and does not award an ordered finishing score.",
  "Sprint participation points are capped at three per driver per season.",
  "A non-points Sprint may be archived for history without affecting championship standings."
 ])
])

article(9,"Championship and Scoring",[
 (1,"Championships",[
  "The MFMA maintains separate Driver, Team, and Vehicle championships.",
  "Only Points Events that are not void contribute to championship standings.",
  "Standings are provisional until an entrant has participated in at least two valid Points Events."
 ]),
 (2,"Regular Points",[
  "At-large event classification awards five points to the winner and three points to the other classified team in the current two-team format.",
  "Each official session awards two points to the session winner and one point to the other classified team. A DNF classification receives one point unless the MRA Steward records a disqualification or no-result outcome.",
  "Driver, Team, and Vehicle standings receive the applicable at-large and session points from the recorded entrant for that result."
 ]),
 (3,"Participation and Drops",[
  "A registered passenger or substitute driver who completes at least one session and is not already the event's scoring driver receives one driver-only participation point.",
  "Passenger participation is limited to one point per event and three points per driver per season. Multiple passenger changes may be recorded, but each eligible passenger is awarded no more than once for the event.",
  "After six valid Points Events, each Driver, Team, and Vehicle standing drops its single lowest at-large event score. Session and participation points are not dropped."
 ]),
 (4,"Adjustments",[
  "The MRA may add or deduct points through an audited adjustment stating the affected championship, entrant, value, reason, official, and date.",
  "Deleting or voiding an adjustment removes its effect without rewriting the underlying event result.",
  "A penalty may affect session results, at-large results, championship points, future participation, or any combination expressly recorded by the MRA Steward."
 ])
])

article(10,"Communications and Digital Race Control",[
 (1,"Official Communications",[
  "Competitors shall maintain a functioning two-way radio unless the MRA Steward announces another approved method.",
  "Physical flags, MRA displays, sound or voice alerts, and Official Radio Channel announcements form one coordinated race-control system.",
  "When a visual flag cannot immediately be observed, the corresponding official radio or MRA digital instruction constitutes notice.",
  "Drivers shall keep radio traffic clear during official communication. False reports, jamming, and intentional interference are serious violations."
 ]),
 (2,"Driver Portal",[
  "The Driver Portal permits team sign-in, approved driver and passenger selection, vehicle reporting, hazard requests, acknowledgements, championship access, and advisory GPS functions.",
  "A Safety Car or Red request from a Driver does not itself activate that final state. Receipt of a new hazard report produces immediate Yellow while the MRA Steward evaluates and resolves the report.",
  "A report marked Pending has not yet been received by MRA. The portal retries the report after connectivity returns.",
  "The MRA Steward remains the only authority permitted to set official flags, time, results, penalties, and event status."
 ]),
 (3,"GPS and Advisory Information",[
  "GPS location, sector proximity, map heading, navigation view, and displayed speed are advisory aids and are not official timing, certified speed measurement, or permission to disregard a signal.",
  "When GPS is poor, stale, unavailable, or outside mapped grounds, all active local warnings remain displayed.",
  "A driver remains responsible for observing conditions and complying with global signals regardless of GPS operation."
 ])
])

article(11,"Conduct Investigations and Penalties",[
 (1,"Standards of Conduct",[
  "Competitors shall act honestly, safely, respectfully, and without intentional interference, contact, blocking, false spotting, outside assistance, communication abuse, or unsportsmanlike conduct.",
  "Competitors shall follow MRA Steward instructions and remain within approved boundaries.",
  "A penalty shall be proportionate to severity, intent, safety effect, competitive effect, and relevant prior conduct."
 ]),
 (2,"Warnings and Investigations",[
  "An Infraction Warning is displayed as White crossed with folded Yellow for ten seconds. It is recorded and competition continues under the applicable Primary Safety State.",
  "White alone opens an investigation, pauses the active session when applicable, and directs competitors to the Starting Zone. White does not itself establish guilt or disqualification.",
  "Every investigation shall receive a case identifier, allegation, affected team or entrant, opening time, status, and final disposition.",
  "The MRA Steward may close an investigation with no action and restore the prior state."
 ]),
 (3,"Penalty Decisions",[
  "Available decisions include no action, warning, time remedy, restart, no result, session disqualification, event disqualification, point adjustment, participation restriction, removal from the event, or removal from MFMA competition.",
  "A decision shall include a rationale and shall be recorded separately from the initial allegation.",
  "A disqualification is announced after the decision. Because competitors are already stopped for review, no separate physical disqualification flag is required.",
  "False or malicious spot reports, intentional radio interference, leaving approved grounds for advantage, and deliberately unsafe conduct ordinarily justify disqualification, subject to the recorded circumstances."
 ])
])

article(12,"Safety Course and Hazard Management",[
 (1,"Driver and Passenger Safety",[
  "Drivers are responsible for vehicle control, speed, visibility, terrain, occupants, competitors, spectators, and property.",
  "Seatbelts shall be used whenever fitted. No person may ride in an unsafe position or exceed safe vehicle occupancy.",
  "A passenger shall not obstruct the driver or engage in conduct that creates a hazard."
 ]),
 (2,"Course and Sectors",[
  "The MRA Steward shall designate legal grounds, closed areas, optional field access, and any building restrictions before competition.",
  "Open areas are legal unless specifically closed. Buildings may be approved hiding locations but are not through routes.",
  "The MRA system divides the grounds into ten core sectors and three optional field zones for local warnings and boundary administration.",
  "The MRA Steward may close or reopen sectors during an event. A closed sector is out of bounds regardless of GPS display."
 ]),
 (3,"Hazards and Conditions",[
  "A competitor aware of an immediate hazard shall report it as soon as practical and shall not conceal it.",
  "The MRA Steward may use Yellow, Safety Car, Red, a local Yellow, a Grip warning, a boundary closure, or an additional digital message according to the risk.",
  "Grip warnings may describe rain or water, loose debris, or another reduced-traction condition and may remain active with Green, Yellow, or Safety Car.",
  "A mechanical failure may require Safety Car, Red, vehicle restriction, replacement approval, restart, or termination."
 ])
])

article(13,"Official Signals",[
 (1,"Authority and Priority",[
  "Only a signal displayed or authorized by the MRA Steward is official. Unauthorized imitation is prohibited.",
  "Primary Safety State precedence is Red, Safety Car, Yellow, then Green. Checkered ends competitive activity but does not override a later Red safety instruction.",
  "Supplemental Grip, local Yellow, Move Over, enforcement, and movement information may coexist with a Primary Safety State but shall never reduce the required safety response."
 ]),
 (2,"Primary Signals",[
  "Green starts or resumes normal competition. The maximum competition speed is 15 mph unless reduced by another instruction.",
  "Yellow indicates caution. Competition continues at no more than 7 mph and the official clock advances at one-half rate.",
  "Safety Car is shown physically by open Yellow with open Red held perpendicular from the Official Vehicle. Overtaking and competitive activity are prohibited.",
  "Red requires competitors to stop safely as soon as practical and await instructions.",
  "Checkered ends the session. Competitors cease competition and follow the next movement order."
 ]),
 (3,"Supplemental and Movement Signals",[
  "Move Over is an MRA digital and radio instruction until a dedicated Blue-with-Yellow-stripe physical flag is acquired.",
  "Grip is an MRA digital and radio instruction until a dedicated striped Yellow-and-Red physical flag is acquired. Ordinary Red shall not be improvised for Grip because Red alone always means stop.",
  "Return to Start is shown by a folded Yellow flag twirled briefly, stopped, and repeated.",
  "Proceed to Line is shown by a folded Green flag twirled briefly, stopped, and then pointed toward the Starting Line.",
  "Qualifying Halfway is shown by Green crossed with Checkered.",
  "The full physical presentation examples in Appendix B are part of the official signal standard."
 ])
])

doc.add_page_break(); title("Appendix A Course and Sector Administration",1)
doc.add_paragraph("The MRA course map is an operating aid. Before each event, the MRA Steward shall confirm that mapped boundaries reflect current legal access, surface conditions, buildings, hazards, and temporary closures.")
add_table(["Label","Operating area","Default status"],[
 ("S1","Starting Zone","Open"),("S2","West Grounds","Open"),("S3","Northwest Grounds","Open"),("S4","North Grounds","Open"),("S5","Northeast Grounds","Open"),("S6","East Loop","Open"),("S7","South Grounds","Open"),("S8N","Central North","Open"),("S8S","Central South","Open"),("S10","Shop Grounds","Open"),("FW","West Field","Closed unless opened"),("FN","North Field","Closed unless opened"),("FE","East Field","Closed unless opened")
], [1.0,3.8,1.8])
doc.add_paragraph("Local signals. Local Yellow, Rain or Water Grip, and Loose Debris Grip may be assigned to one or more open sectors. GPS proximity may emphasize the warning, but the warning remains globally visible when GPS data is limited.")

doc.add_page_break(); title("Appendix B Physical Flag Reference",1)
doc.add_paragraph("Available physical inventory: Green, Yellow, Red, White, and Checkered. Fabric shown as open is fully displayed. Fabric shown as folded is bunched against the pole and held with the pole in the same hand.")
for key in ["green","safety-car","investigation","warning"]:
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_together=True
    p.paragraph_format.space_after=Pt(2); add_accessible_picture(p, diagrams[key], Inches(4.25), f"Official physical signal example: {key.replace('-', ' ')}")
doc.add_page_break(); title("Appendix B Physical Flag Reference Continued",1)
for key in ["return","proceed","halfway"]:
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_together=True
    p.paragraph_format.space_after=Pt(2); add_accessible_picture(p, diagrams[key], Inches(4.25), f"Official physical signal example: {key.replace('-', ' ')}")
doc.add_paragraph("Red safeguard. Red displayed alone always means stop. Ordinary Red shall not be alternated or combined informally to imitate a Grip warning.")
doc.add_paragraph("Digital-only supplements. Move Over uses a Blue-and-Yellow digital display and radio command. Grip uses a Yellow-and-Red digital treatment with the repeated voice callout GRIP GRIP GRIP. These remain supplemental to the active Primary Safety State.")

doc.add_page_break(); title("Appendix C Championship Tables",1)
add_table(["Award","Driver","Team","Vehicle"],[
 ("At-large winner","5","5","5"),("At-large classified","3","3","3"),("Session winner","2","2","2"),("Session classified","1","1","1"),("Eligible DNF","1","1","1"),("Passenger participation","1","0","0"),("Sprint participation","1","0","0"),("Qualifying participation","1","0","0"),("Qualifying fastest bonus","1 additional","0","0")
],[2.25,1.35,1.35,1.35])
add_table(["Policy","Rule"],[
 ("Passenger cap","Maximum 3 passenger participation points per driver per season"),("Sprint cap","Maximum 3 Sprint participation points per driver per season"),("Dropped score","After 6 valid Points Events drop one lowest at-large event score"),("Team affiliation","Team points follow the driver's registered team"),("Non-points","Archived for history but excluded from standings"),("Void event","Retained administratively but excluded from standings"),("Adjustment","Audited additions or deductions; underlying results remain intact")
],[2.0,4.4])

doc.add_page_break(); title("Appendix D Revision Record",1)
add_table(["Edition","Date","Status","Summary"],[
 ("2026 Summer Edition","July 10 2026","Superseded upon adoption","Original comprehensive Sporting Code"),
 ("2026 Final Draft","October 2 2026","Final draft","MRA Steward model, current formats, championship scoring, digital race control, investigation workflow, GPS and sector safety, and illustrated flag standard"),
 ("2026 Final Draft Revision 1","October 3 2026","Final draft","Recorded roll-call governance for mid-season amendments and mid-session suspensions; official team registry corrected")
],[1.35,1.35,1.4,3.0])
doc.add_paragraph("Adoption. Upon formal adoption by the MFMA, this edition supersedes prior rules and practices that conflict with it. Event-specific temporary rulings remain limited to the event in which they are issued unless incorporated into a later adopted edition.")

# Header after content is assembled.
header=sec.header
hp=header.paragraphs[0]; hp.alignment=WD_ALIGN_PARAGRAPH.LEFT
rr=hp.add_run("MFMA   "); rr.bold=True; rr.font.color.rgb=RGBColor.from_string(RED); rr.font.size=Pt(9)
rr=hp.add_run("MRA SPORTING CODE 2026 FINAL DRAFT"); rr.bold=True; rr.font.color.rgb=RGBColor.from_string(BLACK); rr.font.size=Pt(8)

doc.core_properties.title="MFMA Official Sporting Code 2026 Final Draft"
doc.core_properties.subject="Competition rules, championship policy, race control, and official signals"
doc.core_properties.author="Maeschen Farm Motorsports Association"
doc.core_properties.keywords="MFMA, MRA, Sporting Code, Race Control, Motorsport"
doc.save(DOCX)
print(DOCX)
